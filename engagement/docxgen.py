"""DOCX output for an engagement plan (see logic.build_plan)."""

import io

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

NAVY = RGBColor(0x08, 0x3D, 0x5F)
GREY = RGBColor(0x55, 0x65, 0x70)


def _shade(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def _field(run, instr):
    for kind, text in (("begin", None), (None, instr), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)


def _cell_text(cell, text, bold=False, size=9.5, color=None):
    cell.text = ""
    lines = str(text).split("\n")
    p = cell.paragraphs[0]
    for i, line in enumerate(lines):
        if i:
            p = cell.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(line)
        r.bold = bold
        r.font.size = Pt(size)
        if color is not None:
            r.font.color.rgb = color


def _table(doc, header, rows, widths, bold_last=False):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]
        _cell_text(c, h, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
        _shade(c, "083D5F")
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        last = bold_last and ri == len(rows) - 1
        for i, v in enumerate(row):
            _cell_text(cells[i], v, bold=last or (i == 0 and len(header) == 4))
            if last:
                _shade(cells[i], "E8EEF2")
    if widths:
        total = sum(widths)
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Inches(6.5 * w / total)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def make_docx(plan, brief_version):
    f = plan["fields"]
    doc = Document()
    for s in doc.sections:
        s.left_margin = s.right_margin = Inches(1)
        s.top_margin = s.bottom_margin = Inches(0.9)
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)
    st.paragraph_format.space_after = Pt(6)

    sec = doc.sections[0]
    hp = sec.header.paragraphs[0]
    hr = hp.add_run("Verlocity, LLC · Confidential")
    hr.font.size = Pt(8)
    hr.font.color.rgb = GREY
    fp = sec.footer.paragraphs[0]
    fr = fp.add_run(f"Draft for review before issuance · Brief v{brief_version} · Library v{plan['library_version']} · Page ")
    fr.font.size = Pt(8)
    fr.font.color.rgb = GREY
    pr = fp.add_run()
    pr.font.size = Pt(8)
    _field(pr, "PAGE")

    p = doc.add_paragraph()
    r = p.add_run("VERLOCITY, LLC")
    r.bold = True
    r.font.size = Pt(13)
    r.font.color.rgb = NAVY
    p.paragraph_format.space_after = Pt(0)
    p = doc.add_paragraph()
    r = p.add_run("Banker-led growth strategy for community financial institutions")
    r.font.size = Pt(9)
    r.font.color.rgb = GREY
    p = doc.add_paragraph()
    r = p.add_run("STATEMENT OF WORK")
    r.bold = True
    r.font.size = Pt(20)
    r.font.color.rgb = NAVY
    p.paragraph_format.space_after = Pt(0)
    p = doc.add_paragraph()
    r = p.add_run(f"{f['client_name']}: Integrated Deposit Growth Partnership")
    r.bold = True
    r.font.size = Pt(13)
    mods = f["modules"]
    parts = []
    if "assess" in mods:
        parts.append("Deposit Franchise Assessment")
    if any(k in mods for k in ("steps14", "step5", "step6")):
        parts.append("Consumer Campaign")
    parts.append("Strategic Advisory")
    p = doc.add_paragraph()
    r = p.add_run(" · ".join(parts))
    r.font.size = Pt(10)
    r.font.color.rgb = GREY

    h = doc.add_paragraph()
    hr = h.add_run("Engagement Summary")
    hr.bold = True
    hr.font.size = Pt(12)
    hr.font.color.rgb = NAVY
    t = doc.add_table(rows=0, cols=2)
    t.style = "Table Grid"
    for k, v in plan["summary"]:
        cells = t.add_row().cells
        _cell_text(cells[0], k, bold=True)
        _shade(cells[0], "E8EEF2")
        _cell_text(cells[1], v)
        cells[0].width = Inches(1.7)
        cells[1].width = Inches(4.8)
    doc.add_paragraph()

    for s in plan["sections"]:
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(10)
        h.paragraph_format.keep_with_next = True
        hr = h.add_run(s["title"])
        hr.bold = True
        hr.font.size = Pt(12)
        hr.font.color.rgb = NAVY
        for b in s["blocks"]:
            t_ = b["t"]
            if t_ == "p":
                para = doc.add_paragraph()
                run = para.add_run(b["x"])
                if b.get("flag"):
                    run.font.highlight_color = WD_COLOR_INDEX.YELLOW
            elif t_ == "bullet":
                para = doc.add_paragraph(style="List Bullet")
                run = para.add_run(b["x"])
                if b.get("flag"):
                    run.font.highlight_color = WD_COLOR_INDEX.YELLOW
            elif t_ == "h3":
                para = doc.add_paragraph()
                para.paragraph_format.space_before = Pt(8)
                para.paragraph_format.keep_with_next = True
                run = para.add_run(b["x"])
                run.bold = True
                run.font.size = Pt(11)
            elif t_ == "h4":
                para = doc.add_paragraph()
                para.paragraph_format.space_before = Pt(4)
                para.paragraph_format.keep_with_next = True
                run = para.add_run(b["x"])
                run.bold = True
                run.italic = True
            elif t_ == "table":
                _table(doc, b["header"], b["rows"], b.get("widths"), b.get("bold_last", False))
            elif t_ == "sign":
                tbl = doc.add_table(rows=1, cols=2)
                sg = b["signers"]
                for i, (who, side) in enumerate((("client", b["client"]), ("verlocity", "VERLOCITY, LLC"))):
                    c = tbl.rows[0].cells[i]
                    c.text = ""
                    lines = [side, "", "By: ______________________________",
                             f"Name: {sg[who]['name'] or '____________________________'}",
                             f"Title: {sg[who]['title'] or '_____________________________'}",
                             "Date: _____________________________"]
                    for j, line in enumerate(lines):
                        para = c.paragraphs[0] if j == 0 else c.add_paragraph()
                        para.paragraph_format.space_after = Pt(2)
                        run = para.add_run(line)
                        run.bold = j == 0
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
