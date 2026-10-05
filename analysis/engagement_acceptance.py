"""Acceptance test: regenerate a real, hand-written SOW from its facts and compare.

Usage (all three inputs are PRIVATE files kept outside this public repo):
  python analysis/engagement_acceptance.py --library lib.json --fields fields.json --reference ref.docx [--out generated.docx]

  lib.json     library export (clauses, modules, module_text) -- or omit and set
               SUPABASE_URL + SUPABASE_SERVICE_KEY to read the private tables
  fields.json  the reference engagement's facts (client, modules, start, term, ...)
  ref.docx     the hand-written SOW

Hard checks (exit 1 on any miss): summary table, fee table, installment
schedule, month labels. Fixed-clause text is compared paragraph by paragraph;
differences are listed so a reviewer can accept or fix each one.
"""

import argparse
import json
import os
import re
import sys

from docx import Document

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
from engagement import logic  # noqa: E402
from engagement.docxgen import make_docx  # noqa: E402


def norm(s):
    s = (s or "").replace("’", "'").replace("‘", "'").replace("“", '"').replace("”", '"')
    return re.sub(r"\s+", " ", s).strip()


def doc_text(path_or_bytes):
    import io
    d = Document(io.BytesIO(path_or_bytes) if isinstance(path_or_bytes, bytes) else path_or_bytes)
    paras = [norm(p.text) for p in d.paragraphs if norm(p.text)]
    tables = []
    for t in d.tables:
        tables.append([[norm(c.text) for c in row.cells] for row in t.rows])
    return paras, tables


def find_table(tables, header_start):
    for t in tables:
        if t and t[0] and norm(t[0][0]).startswith(header_start):
            return t
    return None


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--library")
    ap.add_argument("--fields", required=True)
    ap.add_argument("--reference", required=True)
    ap.add_argument("--out")
    a = ap.parse_args()

    if a.library:
        library = json.load(open(a.library, encoding="utf-8"))
    else:
        from engagement import store
        library = store.get_library()
    fields = json.load(open(a.fields, encoding="utf-8"))

    plan = logic.build_plan(fields, library)
    print("ready:", plan["ready"])
    for x in plan["assertions"]:
        print(f"  [{x['level']}] {x['text']}")
    if not plan.get("sections"):
        print("plan is not buildable:", plan.get("missing"))
        return 1
    data = make_docx(plan, brief_version=1)
    if a.out:
        open(a.out, "wb").write(data)

    gp, gt = doc_text(data)
    rp, rt = doc_text(a.reference)
    hard_fail = []

    # 1. summary table
    gsum = {r[0]: r[1] for r in gt[0]}
    rsum = {r[0]: r[1] for r in rt[0] if len(r) > 1 and r[0]}
    print("\nSUMMARY TABLE")
    for k in ["Client", "Service Provider", "Term", "Professional Fee", "Included", "Third-Party Costs", "Draft Date", "Status"]:
        g, r = gsum.get(k), rsum.get(k)
        ok = norm(g) == norm(r)
        print(f"  {'OK  ' if ok else 'DIFF'} {k}")
        if not ok:
            print(f"       generated: {g}\n       reference: {r}")
            hard_fail.append("summary:" + k)

    # 2. fee tables
    print("\nFEES")
    gfee, rfee = find_table(gt, "Professional Fee Component"), find_table(rt, "Professional Fee Component")
    gins, rins = find_table(gt, "Installment Schedule"), find_table(rt, "Installment Schedule")
    for name, g, r in (("fee table", gfee, rfee), ("installments", gins, rins)):
        gs = [[norm(c) for c in row] for row in (g or [])]
        rs = [[norm(c) for c in row] for row in (r or [])]
        gs = [row for row in gs if any(row)]
        rs = [row for row in rs if any(row)]
        gs2 = [[c.replace("$", "").replace(",", "") if re.match(r"^\$?[\d,]+$", c) else c for c in row] for row in gs]
        rs2 = [[c.replace("$", "").replace(",", "") if re.match(r"^\$?[\d,]+$", c) else c for c in row] for row in rs]
        ok = gs2 == rs2
        print(f"  {'OK  ' if ok else 'DIFF'} {name}")
        if not ok:
            print("       generated:", gs)
            print("       reference:", rs)
            hard_fail.append(name)

    # 3. month labels in the workplan
    print("\nWORKPLAN MONTHS")
    gwp, rwp = find_table(gt, "Month"), find_table(rt, "Month")
    gm = [re.sub(r"\s+", " ", row[0]) for row in (gwp or [])[1:]]
    rm = [re.sub(r"Month (\d)([A-Z])", r"Month \1 \2", row[0]) for row in (rwp or [])[1:]]
    rm = [re.sub(r"\s+", " ", x) for x in rm]
    ok = gm == rm
    print(f"  {'OK  ' if ok else 'DIFF'} month labels")
    if not ok:
        print("       generated:", gm)
        print("       reference:", rm)
        hard_fail.append("months")

    # 4. fixed clauses, paragraph by paragraph
    print("\nFIXED CLAUSES (generated paragraph found verbatim in the reference?)")
    refset = set(rp) | {c for t in rt for row in t for c in row}
    fixed_titles = ["4. Delivery Approach", "6. Governance and Working Cadence", "9. Assumptions and Dependencies",
                    "10. Out-of-Scope Services", "11. Scope Changes", "12. Review and Acceptance of Deliverables",
                    "13. Term, Termination, and Subsequent Work"]
    fixed_titles.append(f"7. {plan['params']['client']} Responsibilities")
    total = matched = 0
    diffs = []
    for s in plan["sections"]:
        if s["title"] not in fixed_titles and not s["title"].startswith("8."):
            continue
        for b in s["blocks"]:
            if b["t"] in ("p", "bullet"):
                total += 1
                if norm(b["x"]) in refset:
                    matched += 1
                else:
                    diffs.append((s["title"], b["x"]))
    print(f"  {matched} of {total} paragraphs match verbatim")
    for title, x in diffs:
        close = [r for r in rp if norm(x)[:40] == r[:40]]
        print(f"  DIFF in {title}:\n      generated: {norm(x)}")
        if close:
            print(f"      reference: {close[0]}")

    print("\nRESULT:", "PASS (hard checks)" if not hard_fail else "FAIL: " + ", ".join(hard_fail))
    return 1 if hard_fail else 0


if __name__ == "__main__":
    sys.exit(main())
