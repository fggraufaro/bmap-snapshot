"""Tests for the Engagement Builder. Uses a synthetic library (made-up text and
prices) so no real clause language or pricing lives in this public repo.

Run: python -m unittest tests.test_engagement -v
"""

import copy
import io
import json
import os
import sys
import unittest

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
os.environ.setdefault("SUPABASE_SERVICE_KEY", "test")

from docx import Document  # noqa: E402
from flask import Flask  # noqa: E402

from engagement import api, llm, logic  # noqa: E402
from engagement.docxgen import make_docx  # noqa: E402


def lib():
    c = lambda sec, seq, style, body, when=None: {"seq": seq, "style": style, "body": body, "when_modules": when or []}
    clauses = {s: [c(s, 1, "p", f"Text for {s} about {{client}}.")] for s in logic.REQUIRED_SECTIONS}
    clauses["objectives_std"] = [c("o", 1, "bullet", "Std campaign objective.", ["steps14"]),
                                 c("o", 2, "bullet", "Std advisory objective for {client}.", ["advisory"])]
    clauses["assumptions"] = [c("a", 1, "bullet", "Start in {start_month}; {workstream_count_words} workstreams."),
                              c("a", 2, "bullet", "Campaign only.", ["steps14"])]
    clauses["term"] = [c("t", 1, "p", "Term {term_words} months from {start_long} to {end_long}.")]
    clauses["status"] = [c("s", 1, "p", "Draft status")]
    clauses["fees_intro"] = [c("f", 1, "p", "Fees for {fee_scope}.")]
    clauses["fees_intro_campaign"] = [c("f", 1, "p", "Campaign fees for {fee_scope}.")]
    clauses["advisory_note"] = [c("n", 1, "p", "Advisory included in {total}.")]
    clauses["advisory_note_campaign"] = [c("n", 1, "p", "Advisory and campaign included in {total}.")]
    mod = lambda n, price, f, t, dep=None, sort=0: {"name": n, "fee_label": n, "price": price, "from_month": f, "to_month": t,
                                                     "depends_on": dep, "activities": f"{n} work", "tail_activities": None,
                                                     "deliverable": f"{n} output" if n != "Advisory" else None, "sort": sort}
    modules = {"assess": mod("Assess", 3000, 1, 3, None, 10), "steps14": mod("Steps14", 2000, 1, 2, None, 20),
               "step5": mod("Step5", 1000, 2, 3, "steps14", 30), "step6": mod("Step6", 6500, 3, 4, "step5", 40),
               "advisory": mod("Advisory", 0, 1, 99, None, 50)}
    texts = {k: [{"seq": 1, "style": "p", "body": f"Scope of {k} for {{client}} with {{sponsor_phrase}}."}] for k in modules}
    return {"version": 1, "clauses": clauses, "modules": modules, "module_text": texts}


BASE = {"client_name": "Example Bank", "client_short": "Example", "modules": ["assess", "steps14", "step5", "step6"],
        "start": "2027-01-01", "term": 5, "third_party": "pass", "draft_date": "2026-12-01",
        "context": ["New platform"], "objectives": ["Grow deposits"]}


class Logic(unittest.TestCase):
    def test_fees_and_installments(self):
        L = lib()
        plan = logic.build_plan(BASE, L)
        self.assertTrue(plan["ready"])
        f = plan["fees"]
        self.assertEqual(f["total"], 3000 + 2000 + 1000 + 6500)
        self.assertEqual(sum(f["installments"]), f["total"])
        self.assertEqual(len(f["installments"]), 5)
        self.assertEqual(f["base"] % logic.INSTALLMENT_ROUND, 0)

    def test_term_end_dates(self):
        self.assertEqual(str(logic.term_end(logic.dt.date(2026, 9, 1), 5)), "2027-01-31")
        self.assertEqual(str(logic.term_end(logic.dt.date(2027, 12, 1), 3)), "2028-02-29")  # leap year

    def test_term_too_short_blocks(self):
        plan = logic.build_plan(dict(BASE, term=3), lib())
        self.assertFalse(plan["ready"])
        self.assertTrue(any(a["id"] == "term_covers" and a["level"] == "fail" for a in plan["assertions"]))

    def test_dependency_error(self):
        plan = logic.build_plan(dict(BASE, modules=["step5"]), lib())
        self.assertFalse(plan["ready"])
        self.assertTrue(any("requires" in a["text"] for a in plan["assertions"]))

    def test_missing_inputs(self):
        plan = logic.build_plan({"client_name": "X"}, lib())
        self.assertFalse(plan["ready"])
        self.assertIn("services", plan["missing"])

    def test_bad_start_and_term_rejected(self):
        f, errs = logic.clean_fields(dict(BASE, start="2027-01-15", term=40), lib())
        self.assertIsNone(f["start"])
        self.assertIsNone(f["term"])
        self.assertEqual(len(errs), 2)

    def test_unresolved_placeholder_fails(self):
        L = lib()
        L["clauses"]["governance"][0]["body"] = "Oops {not_a_param} here."
        plan = logic.build_plan(BASE, L)
        self.assertFalse(plan["ready"])
        self.assertTrue(any(a["id"] == "placeholders" and a["level"] == "fail" for a in plan["assertions"]))

    def test_incomplete_library_blocks(self):
        L = lib()
        del L["clauses"]["acceptance"]
        self.assertFalse(logic.build_plan(BASE, L)["ready"])

    def test_conditional_clauses(self):
        L = lib()
        p = logic.build_plan(dict(BASE, modules=["assess"], term=3), L)
        texts = [b["x"] for s in p["sections"] for b in s["blocks"] if "x" in b]
        self.assertFalse(any("Campaign only" in t for t in texts))
        p2 = logic.build_plan(BASE, L)
        texts2 = [b["x"] for s in p2["sections"] for b in s["blocks"] if "x" in b]
        self.assertTrue(any("Campaign only" in t for t in texts2))

    def test_narrative_guardrails(self):
        self.assertTrue(logic.check_narrative(["We will add 2 markets."]))
        self.assertTrue(logic.check_narrative(["Verlocity guarantees growth."]))
        self.assertTrue(logic.check_narrative(["Has {placeholder}."]))
        self.assertEqual(logic.check_narrative(["Plain text with additional markets."]), [])

    def test_stale_narrative_warns_and_is_excluded(self):
        L = lib()
        f = dict(BASE, narrative={"purpose": ["Old purpose."], "objectives": [], "key": "stale"})
        plan = logic.build_plan(f, L)
        self.assertTrue(plan["ready"])
        self.assertTrue(any(a["id"] == "narrative" and a["level"] == "warn" for a in plan["assertions"]))
        texts = [b["x"] for s in plan["sections"] for b in s["blocks"] if "x" in b]
        self.assertFalse(any("Old purpose" in t for t in texts))

    def test_fresh_narrative_is_used(self):
        L = lib()
        f, _ = logic.clean_fields(BASE, L)
        f["narrative"] = {"purpose": ["Fresh purpose."], "objectives": ["Grow deposits sensibly."], "key": logic.narrative_key(f)}
        plan = logic.build_plan(f, L)
        texts = [b["x"] for s in plan["sections"] for b in s["blocks"] if "x" in b]
        self.assertTrue(any("Fresh purpose" in t for t in texts))

    def test_docx_builds_and_contains_totals(self):
        L = lib()
        plan = logic.build_plan(BASE, L)
        data = make_docx(plan, 2)
        d = Document(io.BytesIO(data))
        text = "\n".join(p.text for p in d.paragraphs) + "\n".join(c.text for t in d.tables for r in t.rows for c in r.cells)
        self.assertIn("STATEMENT OF WORK", text)
        self.assertIn("Example Bank", text)
        self.assertIn("$12,500", text)


class Llm(unittest.TestCase):
    def test_chat_applies_valid_updates_and_filters_bad_suggestions(self):
        L = lib()
        reply = {"reply": "Which cities?", "updates": {"modules": ["steps14", "step6"], "start": "2027-02-01", "term": 4,
                                                       "context": ["New platform"], "objectives": ["Grow deposits"],
                                                       "third_party": "direct"},
                 "suggestions": [{"field": "term", "value": 99, "why": "bad"},
                                 {"field": "modules", "value": ["assess"], "why": "good"}],
                 "quick_replies": ["Two", "Three"], "flags": []}
        res = llm.chat_turn({"modules": []}, [], "hello", L, complete=lambda p, n: json.dumps(reply))
        self.assertTrue(res["ok"])
        f = res["fields"]
        self.assertEqual(f["modules"], ["steps14", "advisory"])  # step6 dropped: needs step5
        self.assertTrue(res["notes"])
        self.assertEqual((f["start"], f["term"], f["third_party"]), ("2027-02-01", 4, "direct"))
        self.assertEqual([s["field"] for s in res["suggestions"]], ["modules"])
        self.assertEqual(res["suggestions"][0]["label"], "Assess")

    def test_chat_prompt_hides_prices(self):
        seen = {}

        def comp(p, n):
            seen["p"] = p
            return json.dumps({"reply": "ok", "updates": {}, "suggestions": [], "quick_replies": [], "flags": []})
        llm.chat_turn({}, [{"w": "me", "t": "hi"}], "next", lib(), complete=comp)
        self.assertNotIn("6500", seen["p"])
        self.assertIn("steps14", seen["p"])
        self.assertIn("CMO: hi", seen["p"])

    def test_chat_bad_json_does_not_raise(self):
        res = llm.chat_turn({"modules": ["assess"]}, [], "x", lib(), complete=lambda p, n: "not json at all")
        self.assertFalse(res["ok"])
        self.assertEqual(res["fields"], {"modules": ["assess"]})

    def test_chat_fenced_json_is_parsed(self):
        raw = "```json\n" + json.dumps({"reply": "hi", "updates": {}, "suggestions": [], "quick_replies": [], "flags": []}) + "\n```"
        self.assertTrue(llm.chat_turn({}, [], "x", lib(), complete=lambda p, n: raw)["ok"])

    def test_narrative_retries_then_succeeds(self):
        calls = []

        def comp(p, n):
            calls.append(p)
            if len(calls) == 1:
                return json.dumps({"purpose": ["We will open 2 cities."], "objectives": ["Grow."]})
            return json.dumps({"purpose": ["We will explore additional cities."], "objectives": ["Grow deposits."]})
        nar = llm.draft_narrative({"client_short": "X", "context": ["a"], "objectives": ["b"]}, complete=comp)
        self.assertEqual(len(calls), 2)
        self.assertIn("failed these checks", calls[1])
        self.assertEqual(logic.check_narrative(nar["purpose"] + nar["objectives"]), [])

    def test_narrative_fails_closed(self):
        with self.assertRaises(llm.NarrativeError):
            llm.draft_narrative({"client_short": "X", "context": ["a"]},
                                complete=lambda p, n: json.dumps({"purpose": ["Guaranteed 5 percent."], "objectives": []}))

    def test_narrative_needs_inputs(self):
        with self.assertRaises(llm.NarrativeError):
            llm.draft_narrative({"client_short": "X"}, complete=lambda p, n: "{}")

    def test_missing_key_is_unavailable(self):
        old = llm.ANTH_KEY
        llm.ANTH_KEY = ""
        try:
            with self.assertRaises(llm.LLMUnavailable):
                llm.default_complete("hi")
        finally:
            llm.ANTH_KEY = old


class FakeStore:
    class StoreError(Exception):
        pass

    def __init__(self):
        self.lib = lib()
        self.engs, self.vers, self.drafts, self.events = {}, {}, {}, []
        self.n = 0

    def get_library(self):
        return copy.deepcopy(self.lib)

    def log_event(self, eid, event, actor=None, detail=None):
        self.events.append((eid, event, actor, detail))

    def create_engagement(self, name, fields, actor=None):
        eid = f"e{len(self.engs) + 1}"
        self.engs[eid] = {"id": eid, "client_name": name, "status": "draft", "current_version": 1,
                          "approved_version": None, "approved_by": None}
        self.vers[(eid, 1)] = {"fields": fields, "transcript": [], "change_log": []}
        return self.engs[eid]

    def list_engagements(self):
        return list(self.engs.values())

    def get_engagement(self, eid):
        if eid not in self.engs:
            return None
        e = self.engs[eid]
        return {"engagement": copy.deepcopy(e), "version": copy.deepcopy(self.vers[(eid, e["current_version"])]),
                "drafts": [dict(d, docx_b64=None) for d in self.drafts.values() if d["engagement_id"] == eid]}

    def save_fields(self, eid, fields, transcript, change_log, changed):
        e = self.engs[eid]
        if changed:
            e["current_version"] += 1
        self.vers[(eid, e["current_version"])] = {"fields": copy.deepcopy(fields), "transcript": transcript, "change_log": change_log}
        return e["current_version"]

    def set_current_fields(self, eid, version, fields, transcript):
        self.vers[(eid, version)] = dict(self.vers[(eid, version)], fields=copy.deepcopy(fields), transcript=transcript)

    def add_draft(self, eid, bv, lv, assertions, narrative, b64, sha):
        self.n += 1
        d = {"id": f"d{self.n}", "engagement_id": eid, "brief_version": bv, "library_version": lv, "status": "draft",
             "assertions": assertions, "docx_b64": b64, "sha256": sha, "reviewed_by": None}
        self.drafts[d["id"]] = d
        return {k: v for k, v in d.items() if k != "docx_b64"}

    def get_draft(self, did, with_docx=False):
        d = self.drafts.get(did)
        if not d:
            return None
        return dict(d) if with_docx else {k: v for k, v in d.items() if k != "docx_b64"}

    def set_draft_status(self, did, status, actor, note=None):
        self.drafts[did]["status"] = status
        if status == "reviewed":
            self.drafts[did]["reviewed_by"] = actor
        return {k: v for k, v in self.drafts[did].items() if k != "docx_b64"}

    def supersede_other_drafts(self, eid, keep):
        for d in self.drafts.values():
            if d["engagement_id"] == eid and d["id"] != keep and d["status"] in ("draft", "reviewed", "sent"):
                d["status"] = "superseded"

    def set_engagement_status(self, eid, status, approved_version=None, approved_by=None):
        e = self.engs[eid]
        e["status"] = status
        if status == "approved":
            e["approved_version"], e["approved_by"] = approved_version, approved_by
        return e


class Api(unittest.TestCase):
    def setUp(self):
        self.fake = FakeStore()
        api.store = self.fake
        self.replies = []
        api.complete = lambda p, n: self.replies.pop(0)
        app = Flask(__name__)
        app.config["TESTING"] = True

        def require_session(fn):
            from functools import wraps
            from flask import request, jsonify

            @wraps(fn)
            def w(*a, **k):
                if request.method != "OPTIONS" and request.headers.get("Authorization") != "Bearer ok":
                    return jsonify({"error": "unauthorized"}), 401
                return fn(*a, **k)
            return w
        api.register(app, require_session)
        self.c = app.test_client()
        self.h = {"Authorization": "Bearer ok"}

    def post(self, url, body=None, h=True):
        return self.c.post(url, json=body or {}, headers=self.h if h else {})

    def make(self):
        return self.post("/engagement/create", {"client_name": "Example Bank", "client_short": "Example"}).get_json()["id"]

    def fill(self, eid, **over):
        upd = {"modules": ["assess", "steps14", "step5", "step6"], "start": "2027-01-01", "term": 5, "third_party": "pass",
               "context": ["New platform"], "objectives": ["Grow deposits"]}
        upd.update(over)
        return self.post(f"/engagement/{eid}/fields", {"updates": upd})

    def test_requires_login(self):
        self.assertEqual(self.c.get("/engagement/list").status_code, 401)
        self.assertEqual(self.post("/engagement/create", {"client_name": "X"}, h=False).status_code, 401)

    def test_full_flow_and_signoff(self):
        eid = self.make()
        r = self.fill(eid)
        self.assertEqual(r.status_code, 200)
        self.assertTrue(r.get_json()["preview"]["ready"])
        # generate before narrative is allowed (warn only) and carries placeholders
        g = self.post(f"/engagement/{eid}/generate")
        self.assertEqual(g.status_code, 200)
        did = g.get_json()["draft"]["id"]
        dl = self.c.get(f"/engagement/{eid}/draft/{did}/docx", headers=self.h)
        self.assertEqual(dl.status_code, 200)
        self.assertTrue(dl.data.startswith(b"PK"))
        # transitions
        s = f"/engagement/{eid}/draft/{did}/status"
        self.assertEqual(self.post(s, {"status": "approved", "actor": "A"}).status_code, 409)  # cannot skip
        self.assertEqual(self.post(s, {"status": "reviewed"}).status_code, 400)  # needs a name
        self.assertEqual(self.post(s, {"status": "reviewed", "actor": "Francisco"}).status_code, 200)
        self.assertEqual(self.post(s, {"status": "sent", "actor": "Francisco"}).status_code, 200)
        self.assertEqual(self.post(s, {"status": "approved", "actor": "Francisco", "note": "signed pdf"}).status_code, 200)
        e = self.fake.engs[eid]
        self.assertEqual((e["status"], e["approved_version"], e["approved_by"]), ("approved", e["current_version"], "Francisco"))
        # changing the brief reopens it and blocks the old draft
        self.fill(eid, term=6)
        self.assertEqual(self.fake.engs[eid]["status"], "draft")
        self.assertTrue(any(ev[1] == "reopened" for ev in self.fake.events))

    def test_generate_blocked_when_not_ready(self):
        eid = self.make()
        g = self.post(f"/engagement/{eid}/generate")
        self.assertEqual(g.status_code, 422)
        self.assertIn("services", g.get_json()["missing"])
        self.fill(eid, term=3)  # step6 needs 4 months
        g = self.post(f"/engagement/{eid}/generate")
        self.assertEqual(g.status_code, 422)

    def test_approval_blocked_after_brief_change(self):
        eid = self.make()
        self.fill(eid)
        did = self.post(f"/engagement/{eid}/generate").get_json()["draft"]["id"]
        self.fill(eid, objectives=["Something else"])
        r = self.post(f"/engagement/{eid}/draft/{did}/status", {"status": "reviewed", "actor": "A"})
        self.assertEqual(r.status_code, 409)

    def test_chat_persists_and_returns_suggestions(self):
        eid = self.make()
        self.replies.append(json.dumps({"reply": "Which cities?", "updates": {"context": ["New platform"], "term": 5},
                                        "suggestions": [{"field": "modules", "value": ["assess"], "why": "Fits."}],
                                        "quick_replies": ["Two"], "flags": []}))
        r = self.post(f"/engagement/{eid}/chat", {"message": "We have a new platform."}).get_json()
        self.assertEqual(r["fields"]["term"], 5)
        self.assertEqual(r["suggestions"][0]["label"], "Assess")
        got = self.c.get(f"/engagement/{eid}", headers=self.h).get_json()
        self.assertEqual([t["w"] for t in got["transcript"]], ["me", "ai"])
        acc = self.post(f"/engagement/{eid}/suggestion", {"field": "modules", "value": ["assess"]}).get_json()
        self.assertEqual(acc["fields"]["modules"], ["assess", "advisory"])
        bad = self.post(f"/engagement/{eid}/suggestion", {"field": "term", "value": 99})
        self.assertEqual(bad.status_code, 400)

    def test_narrative_route_and_unavailable_key(self):
        eid = self.make()
        self.fill(eid)
        self.replies.append(json.dumps({"purpose": ["Example seeks growth."], "objectives": ["Grow deposits."]}))
        r = self.post(f"/engagement/{eid}/narrative")
        self.assertEqual(r.status_code, 200)
        self.assertTrue(r.get_json()["preview"]["ready"])
        pv = r.get_json()["preview"]
        self.assertTrue(any(a["id"] == "narrative" and a["level"] == "pass" for a in pv["assertions"]))
        # a failed draft is rejected, not saved
        self.replies.extend([json.dumps({"purpose": ["Adds 2 markets."], "objectives": []})] * 2)
        self.assertEqual(self.post(f"/engagement/{eid}/narrative").status_code, 422)

        def boom(p, n):
            raise llm.LLMUnavailable("no key")
        api.complete = boom
        self.assertEqual(self.post(f"/engagement/{eid}/narrative").status_code, 503)


if __name__ == "__main__":
    unittest.main()
